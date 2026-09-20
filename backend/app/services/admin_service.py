"""
Admin service — business logic for platform-wide admin operations.
Supports filtering and multi-format exports (Excel, PDF, Word).
"""
import io
from typing import Dict, Any, List
from sqlalchemy.ext.asyncio import AsyncSession
from fastapi import HTTPException, status

from app.repositories.admin_repository import AdminRepository


def _serialize_user(u, app_count: int = 0) -> dict:
    return {
        "id": u.id,
        "email": u.email,
        "first_name": u.first_name or "",
        "last_name": u.last_name or "",
        "role": u.role or "user",
        "is_active": getattr(u, "is_active", True),
        "title": u.title or "",
        "location": u.location or "",
        "created_at": u.created_at.isoformat() if u.created_at else None,
        "application_count": app_count,
    }


def _serialize_job(j) -> dict:
    return {
        "id": j.id,
        "title": j.title,
        "company_name": j.company_name,
        "mode": j.mode,
        "level": j.level,
        "employment_type": j.employment_type,
        "salary_min": j.salary_min,
        "salary_max": j.salary_max,
        "currency": j.currency,
        "is_hot": j.is_hot,
        "is_expired": getattr(j, "is_expired", False),
        "source_url": j.source_url,
        "match_score": j.match_score,
        "ats_score": j.ats_score,
        "deadline": j.deadline.isoformat() if getattr(j, "deadline", None) else None,
        "created_at": j.created_at.isoformat() if getattr(j, "created_at", None) else None,
    }


def _serialize_company(c) -> dict:
    return {
        "id": c.id,
        "name": c.name,
        "sector": c.sector,
        "location": c.location,
        "tier": c.tier,
        "open_roles_count": c.open_roles_count,
        "career_url": c.career_url,
        "contact_email": c.contact_email,
        "ats_platform": c.ats_platform,
        "created_at": c.created_at.isoformat() if getattr(c, "created_at", None) else None,
    }


class AdminService:
    def __init__(self, db: AsyncSession):
        self.repo = AdminRepository(db)

    # ─── Users ────────────────────────────────────────────────────────────────
    async def list_users(self, skip: int, limit: int, q: str, role: str = "", active_status: str = "") -> list[dict]:
        users = await self.repo.list_users(skip, limit, q, role, active_status)
        result = []
        for u in users:
            app_count = await self.repo.user_application_count(u.id)
            result.append(_serialize_user(u, app_count))
        return result

    async def list_users_paginated(self, skip: int, limit: int, q: str, role: str = "", active_status: str = "") -> dict:
        users = await self.repo.list_users(skip, limit, q, role, active_status)
        total = await self.repo.count_users_filtered(q, role, active_status)
        result = []
        for u in users:
            app_count = await self.repo.user_application_count(u.id)
            result.append(_serialize_user(u, app_count))
        return {
            "users": result,
            "total": total,
            "skip": skip,
            "limit": limit,
            "has_more": (skip + limit) < total,
        }

    async def create_user(self, payload: dict) -> dict:
        email = payload.get("email", "").strip().lower()
        if not email:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Email address is required.")
        existing = await self.repo.get_user_by_email(email)
        if existing:
            raise HTTPException(status.HTTP_409_CONFLICT, f"An account with email '{email}' already exists.")
        
        password = payload.pop("password", "")
        if len(password) < 8:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Password must be at least 8 characters.")
        
        from app.core.security import hash_password
        payload["email"] = email
        payload["password_hash"] = hash_password(password)
        
        u = await self.repo.create_user(payload)
        app_count = await self.repo.user_application_count(u.id)
        return _serialize_user(u, app_count)

    async def get_user(self, user_id: int) -> dict:
        u = await self.repo.get_user(user_id)
        if not u:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
        app_count = await self.repo.user_application_count(user_id)
        return _serialize_user(u, app_count)

    async def update_user_role(self, user_id: int, role: str, admin_id: int | None = None) -> dict:
        if role not in ("user", "admin"):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Role must be 'user' or 'admin'")
        if admin_id and admin_id == user_id and role != "admin":
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "You cannot demote your own admin account.")
        u = await self.repo.update_user(user_id, {"role": role})
        if not u:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
        app_count = await self.repo.user_application_count(user_id)
        return _serialize_user(u, app_count)

    async def suspend_user(self, user_id: int, suspend: bool, admin_id: int | None = None) -> dict:
        if admin_id and admin_id == user_id:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "You cannot suspend your own admin account.")
        u = await self.repo.update_user(user_id, {"is_active": not suspend})
        if not u:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
        app_count = await self.repo.user_application_count(user_id)
        return _serialize_user(u, app_count)

    async def update_user_profile(self, user_id: int, payload: dict) -> dict:
        """Admin-level edit of any user's profile fields."""
        ALLOWED_FIELDS = {
            "first_name", "last_name", "email", "title", "location",
            "phone", "role", "is_active", "years_experience"
        }
        patch = {k: v for k, v in payload.items() if k in ALLOWED_FIELDS and v is not None}
        if not patch:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "No valid fields to update.")
        if "email" in patch:
            patch["email"] = patch["email"].strip().lower()
            existing = await self.repo.get_user_by_email(patch["email"])
            if existing and existing.id != user_id:
                raise HTTPException(status.HTTP_409_CONFLICT, "Email already in use by another account.")
        u = await self.repo.update_user(user_id, patch)
        if not u:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
        app_count = await self.repo.user_application_count(user_id)
        return _serialize_user(u, app_count)

    async def reset_user_password(self, user_id: int, new_password: str) -> dict:
        """Admin: force-reset any user's password."""
        if len(new_password) < 8:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Password must be at least 8 characters.")
        from app.core.security import hash_password
        hashed = hash_password(new_password)
        u = await self.repo.update_user(user_id, {"password_hash": hashed})
        if not u:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
        return {"success": True, "message": f"Password for user {user_id} reset successfully."}

    async def delete_user(self, user_id: int, admin_id: int | None = None) -> bool:
        if admin_id and admin_id == user_id:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "You cannot delete your own admin account.")
        deleted = await self.repo.delete_user(user_id)
        if not deleted:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
        return True

    async def promote_by_email(self, email: str) -> dict:
        clean_email = email.strip().lower()
        u = await self.repo.get_user_by_email(clean_email)
        if not u:
            raise HTTPException(status.HTTP_404_NOT_FOUND, f"No user found with email: {email}")
        updated = await self.repo.update_user(u.id, {"role": "admin"})
        if not updated:
            raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "Failed to promote user.")
        app_count = await self.repo.user_application_count(u.id)
        return _serialize_user(updated, app_count)

    # ─── Analytics ────────────────────────────────────────────────────────────
    async def platform_analytics(self) -> dict:
        return await self.repo.platform_analytics()

    # ─── Jobs ─────────────────────────────────────────────────────────────────
    async def list_all_jobs(
        self, skip: int, limit: int, mode: str = "", level: str = "", is_expired: str = ""
    ) -> list[dict]:
        jobs = await self.repo.list_all_jobs(skip, limit, mode, level, is_expired)
        return [_serialize_job(j) for j in jobs]

    async def delete_job(self, job_id: int) -> bool:
        deleted = await self.repo.delete_job(job_id)
        if not deleted:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Job not found")
        return True

    # ─── Companies ────────────────────────────────────────────────────────────
    async def list_all_companies(self, skip: int, limit: int, tier: str = "", sector: str = "") -> list[dict]:
        companies = await self.repo.list_all_companies(skip, limit, tier, sector)
        return [_serialize_company(c) for c in companies]

    async def delete_company(self, company_id: int) -> bool:
        deleted = await self.repo.delete_company(company_id)
        if not deleted:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Company not found")
        return True

    # ─── Export Engine (Excel, PDF, Word) ──────────────────────────────────────
    async def export_data(self, resource: str, format_type: str) -> tuple[bytes, str, str]:
        """Returns (file_bytes, media_type, filename)."""
        data: List[dict] = []
        title = ""
        headers: List[str] = []

        if resource == "users":
            data = await self.list_users(0, 1000, "")
            title = "Platform Users Report"
            headers = ["ID", "Email", "First Name", "Last Name", "Role", "Status", "Title", "Applications", "Joined"]
        elif resource == "jobs":
            data = await self.list_all_jobs(0, 1000)
            title = "Platform Jobs Listings Report"
            headers = ["ID", "Job Title", "Company", "Mode", "Level", "Type", "Salary", "Match Score", "Expired"]
        elif resource == "companies":
            data = await self.list_all_companies(0, 1000)
            title = "Platform Companies Tracked Report"
            headers = ["ID", "Company Name", "Sector", "Location", "Tier", "Open Roles", "ATS Platform", "Contact Email"]
        else:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid resource requested")

        fmt = format_type.lower()
        if fmt in ["excel", "xlsx"]:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = resource.capitalize()
            ws.append([title])
            ws.append([])
            ws.append(headers)

            for item in data:
                if resource == "users":
                    ws.append([item["id"], item["email"], item["first_name"], item["last_name"], item["role"], "Active" if item["is_active"] else "Suspended", item["title"], item["application_count"], item["created_at"]])
                elif resource == "jobs":
                    ws.append([item["id"], item["title"], item["company_name"], item["mode"], item["level"], item["employment_type"], f"{item['currency']} {item['salary_min']}-{item['salary_max']}", item["match_score"], "Yes" if item.get("is_expired") else "No"])
                elif resource == "companies":
                    ws.append([item["id"], item["name"], item["sector"], item["location"], item["tier"], item["open_roles_count"], item["ats_platform"], item["contact_email"]])

            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"admin_{resource}.xlsx"

        elif fmt == "pdf":
            from reportlab.lib.pagesizes import letter, landscape
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors

            buf = io.BytesIO()
            doc = SimpleDocTemplate(buf, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
            story = []
            styles = getSampleStyleSheet()

            story.append(Paragraph(f"<b>{title}</b>", styles["Heading1"]))
            story.append(Spacer(1, 12))

            table_rows = [headers]
            for item in data:
                if resource == "users":
                    table_rows.append([str(item["id"]), item["email"][:25], item["first_name"], item["last_name"], item["role"], "Active" if item["is_active"] else "Suspended", item["title"][:20], str(item["application_count"]), str(item["created_at"])[:10]])
                elif resource == "jobs":
                    table_rows.append([str(item["id"]), item["title"][:25], item["company_name"][:20], item["mode"], item["level"], item["employment_type"], f"{item['salary_min']}-{item['salary_max']}", f"{item['match_score']}%", "Yes" if item.get("is_expired") else "No"])
                elif resource == "companies":
                    table_rows.append([str(item["id"]), item["name"][:25], item["sector"][:20], item["location"][:15], f"Tier {item['tier']}", str(item["open_roles_count"]), item["ats_platform"][:15], item["contact_email"][:20]])

            t = Table(table_rows)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('BOTTOMPADDING', (0,0), (-1,0), 6),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ]))
            story.append(t)
            doc.build(story)
            return buf.getvalue(), "application/pdf", f"admin_{resource}.pdf"

        elif fmt in ["word", "docx"]:
            import docx
            doc = docx.Document()
            doc.add_heading(title, 0)

            table = doc.add_table(rows=1, cols=len(headers))
            hdr_cells = table.rows[0].cells
            for i, h in enumerate(headers):
                hdr_cells[i].text = h

            for item in data:
                row_cells = table.add_row().cells
                if resource == "users":
                    vals = [str(item["id"]), item["email"], item["first_name"], item["last_name"], item["role"], "Active" if item["is_active"] else "Suspended", item["title"], str(item["application_count"]), str(item["created_at"])[:10]]
                elif resource == "jobs":
                    vals = [str(item["id"]), item["title"], item["company_name"], item["mode"], item["level"], item["employment_type"], f"{item['salary_min']}-{item['salary_max']}", f"{item['match_score']}%", "Yes" if item.get("is_expired") else "No"]
                elif resource == "companies":
                    vals = [str(item["id"]), item["name"], item["sector"], item["location"], f"Tier {item['tier']}", str(item["open_roles_count"]), item["ats_platform"], item["contact_email"]]
                
                for idx, v in enumerate(vals):
                    row_cells[idx].text = v

            buf = io.BytesIO()
            doc.save(buf)
            return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"admin_{resource}.docx"

        else:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Format must be excel, pdf, or word")

    # ─── System Email & SMTP Settings ──────────────────────────────────────────
    async def get_email_settings(self) -> dict:
        from app.core.config import settings
        return {
            "smtp_host": settings.smtp_host,
            "smtp_port": settings.smtp_port,
            "smtp_user": settings.smtp_user,
            "smtp_password_set": bool(settings.smtp_password),
            "smtp_from_email": settings.smtp_from_email,
            "smtp_from_name": settings.smtp_from_name,
            "emails_enabled": settings.emails_enabled,
        }

    async def update_email_settings(self, payload: dict) -> dict:
        from app.core.config import settings
        if "smtp_host" in payload and payload["smtp_host"] is not None:
            settings.smtp_host = str(payload["smtp_host"])
        if "smtp_port" in payload and payload["smtp_port"] is not None:
            settings.smtp_port = int(payload["smtp_port"])
        if "smtp_user" in payload and payload["smtp_user"] is not None:
            settings.smtp_user = str(payload["smtp_user"])
        if "smtp_password" in payload and payload["smtp_password"]:
            settings.smtp_password = str(payload["smtp_password"])
        if "smtp_from_email" in payload and payload["smtp_from_email"] is not None:
            settings.smtp_from_email = str(payload["smtp_from_email"])
        if "smtp_from_name" in payload and payload["smtp_from_name"] is not None:
            settings.smtp_from_name = str(payload["smtp_from_name"])
        if "emails_enabled" in payload and payload["emails_enabled"] is not None:
            settings.emails_enabled = bool(payload["emails_enabled"])
        return await self.get_email_settings()

    async def test_smtp_connection(self, target_email: str | None = None) -> dict:
        """
        Test SMTP connectivity with full diagnostic output.
        Auto-detects SSL (port 465) vs STARTTLS (port 587).
        Returns the full error detail so it can be shown in the admin UI.
        """
        import smtplib
        import ssl
        from email.mime.multipart import MIMEMultipart
        from email.mime.text import MIMEText as MIMEPlain
        from app.core.config import settings

        if not settings.smtp_host or not settings.smtp_port:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="SMTP host and port must be configured before testing connection."
            )

        if not settings.smtp_user:
            return {
                "success": False,
                "message": "SMTP_USER is not set. Add it in Render dashboard → Environment.",
                "config": {"host": settings.smtp_host, "port": settings.smtp_port, "user": "<not set>"}
            }

        if not settings.smtp_password:
            return {
                "success": False,
                "message": "SMTP_PASSWORD is not set. Add Gmail App Password in Render dashboard → Environment.",
                "config": {"host": settings.smtp_host, "port": settings.smtp_port, "user": settings.smtp_user}
            }

        # For Gmail, From MUST match authenticated user
        is_gmail = "gmail" in (settings.smtp_host or "").lower()
        from_email = settings.smtp_user if is_gmail else (settings.smtp_from_email or settings.smtp_user)
        from_name = settings.smtp_from_name or "Denno Career OS"
        recipient = target_email or settings.smtp_user  # Default: send test to yourself

        # Strip spaces from App Password (common copy-paste issue)
        smtp_password = (settings.smtp_password or "").replace(" ", "")

        # Auto-detect SSL from port
        port = settings.smtp_port
        use_ssl = port == 465

        msg = MIMEMultipart("alternative")
        msg["Subject"] = "[Denno] SMTP Test — Connection Successful"
        msg["From"] = f"{from_name} <{from_email}>"
        msg["To"] = recipient
        body_text = (
            f"This is an automated SMTP diagnostic test from Denno Career OS.\n\n"
            f"Config: {from_email} via {settings.smtp_host}:{port} "
            f"({'SSL' if use_ssl else 'STARTTLS'})\n\n"
            f"If you receive this, password reset emails will work correctly."
        )
        msg.attach(MIMEPlain(body_text, "plain", "utf-8"))

        try:
            ctx = ssl.create_default_context()
            if use_ssl:
                # Port 465 — direct SSL
                with smtplib.SMTP_SSL(settings.smtp_host, port, timeout=15, context=ctx) as server:
                    server.login(settings.smtp_user, smtp_password)
                    server.sendmail(from_email, [recipient], msg.as_string())
            else:
                # Port 587 — STARTTLS
                with smtplib.SMTP(settings.smtp_host, port, timeout=15) as server:
                    server.ehlo()
                    server.starttls(context=ctx)
                    server.ehlo()
                    server.login(settings.smtp_user, smtp_password)
                    server.sendmail(from_email, [recipient], msg.as_string())

            return {
                "success": True,
                "message": f"✓ Test email sent successfully to {recipient}. Check the inbox!",
                "recipient": recipient,
                "config": {
                    "host": settings.smtp_host,
                    "port": port,
                    "mode": "SSL" if use_ssl else "STARTTLS",
                    "user": settings.smtp_user,
                    "from": from_email,
                }
            }
        except smtplib.SMTPAuthenticationError as exc:
            return {
                "success": False,
                "error_type": "AUTH_FAILED",
                "message": (
                    f"Gmail rejected the App Password. "
                    f"Make sure you are using a 16-character App Password (not your login password). "
                    f"Generate one at: https://myaccount.google.com/apppasswords — error: {exc}"
                ),
                "config": {"host": settings.smtp_host, "port": port, "user": settings.smtp_user}
            }
        except smtplib.SMTPException as exc:
            return {
                "success": False,
                "error_type": "SMTP_ERROR",
                "message": f"SMTP error: {exc}",
                "config": {"host": settings.smtp_host, "port": port, "user": settings.smtp_user}
            }
        except OSError as exc:
            return {
                "success": False,
                "error_type": "CONNECTION_ERROR",
                "message": (
                    f"Could not connect to {settings.smtp_host}:{port}. "
                    f"Check SMTP_HOST and SMTP_PORT settings. Error: {exc}"
                ),
                "config": {"host": settings.smtp_host, "port": port, "user": settings.smtp_user}
            }
        except Exception as exc:
            return {
                "success": False,
                "error_type": "UNKNOWN",
                "message": f"Unexpected error: {type(exc).__name__}: {exc}",
                "config": {"host": settings.smtp_host, "port": port, "user": settings.smtp_user}
            }

